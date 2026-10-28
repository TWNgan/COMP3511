import sys, site
sys.path.append(site.getusersitepackages())
from pptx import Presentation
prs = Presentation(sys.argv[1])
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            t = shp.text_frame.text.strip()
            if t:
                texts.append(t)
    print(f"--- slide {i} ---")
    print("\n".join(texts))
