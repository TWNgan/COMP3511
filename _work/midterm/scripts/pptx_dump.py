"""Dump text (incl. tables, groups, notes) of .pptx files: pptx_dump.py out_dir file1.pptx [file2 ...]"""
import os
import site
import sys

sys.path.append(site.getusersitepackages())
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")


def shape_texts(shp, out, pics):
    if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shp.shapes:
            shape_texts(s, out, pics)
        return
    if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
        pics.append(shp)
    if getattr(shp, "has_table", False) and shp.has_table:
        for row in shp.table.rows:
            out.append(" | ".join(c.text.strip().replace("\n", " / ") for c in row.cells))
    if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
        t = shp.text_frame.text.strip()
        if t:
            out.append(t)


out_dir = sys.argv[1]
for path in sys.argv[2:]:
    prs = Presentation(path)
    name = os.path.splitext(os.path.basename(path))[0].replace(" ", "_").replace("(", "").replace(")", "")
    lines, npics = [], 0
    for i, slide in enumerate(prs.slides, 1):
        out, pics = [], []
        for shp in slide.shapes:
            shape_texts(shp, out, pics)
        npics += len(pics)
        lines.append("===== SLIDE %d =====%s" % (i, ("  [%d picture(s)]" % len(pics)) if pics else ""))
        lines.extend(out)
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                lines.append("[notes] " + n)
    with open(os.path.join(out_dir, name + ".txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("%-40s slides=%3d pictures=%3d chars=%6d" % (name, len(prs.slides), npics, sum(len(l) for l in lines)))
