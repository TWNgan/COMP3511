"""Extract picture shapes from given slides: pptx_images.py file.pptx out_dir [slide numbers...] (all if none)"""
import os
import site
import sys

sys.path.append(site.getusersitepackages())
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402


def pics(shapes):
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from pics(s.shapes)
        elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield s


path, out_dir = sys.argv[1], sys.argv[2]
want = set(int(x) for x in sys.argv[3:])
os.makedirs(out_dir, exist_ok=True)
prs = Presentation(path)
stem = os.path.splitext(os.path.basename(path))[0].replace(" ", "_").replace("(", "").replace(")", "")
for i, slide in enumerate(prs.slides, 1):
    if want and i not in want:
        continue
    for k, p in enumerate(pics(slide.shapes), 1):
        img = p.image
        name = "%s_s%02d_%d.%s" % (stem, i, k, img.ext)
        with open(os.path.join(out_dir, name), "wb") as f:
            f.write(img.blob)
        print(name, img.size if hasattr(img, "size") else "", len(img.blob))
